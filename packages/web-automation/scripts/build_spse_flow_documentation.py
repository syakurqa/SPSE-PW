from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "SPSE_Flow_1_2_3_Execution_Guide.docx"
IMG1 = ROOT / "docs" / "screenshots" / "flow3-pokja-package-list.png"
IMG2 = ROOT / "docs" / "screenshots" / "flow3-package-complete.png"


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_table(doc, rows, widths):
    table = doc.add_table(rows=1, cols=len(widths))
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    table.style = "Table Grid"
    for i, text in enumerate(rows[0]):
        cell = table.rows[0].cells[i]
        cell.width = Inches(widths[i])
        cell.text = text
        shade(cell, "E8EEF5")
        set_cell_margins(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row_data in rows[1:]:
        cells = table.add_row().cells
        for i, text in enumerate(row_data):
            cells[i].width = Inches(widths[i])
            cells[i].text = text
            set_cell_margins(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return table


def add_step(doc, title, detail):
    p = doc.add_paragraph(style="List Number")
    r = p.add_run(title + ". ")
    r.bold = True
    p.add_run(detail)


doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(1)
sec.bottom_margin = Inches(1)
sec.left_margin = Inches(1)
sec.right_margin = Inches(1)
sec.header_distance = Inches(0.492)
sec.footer_distance = Inches(0.492)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25
for name, size, before, after, color in (
    ("Heading 1", 16, 18, 10, "2E74B5"),
    ("Heading 2", 13, 14, 7, "2E74B5"),
    ("Heading 3", 12, 10, 5, "1F4D78"),
):
    style = styles[name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("Dokumentasi Eksekusi SPSE\nFlow 1, Flow 2, dan Flow 3")
r.bold = True
r.font.size = Pt(24)
r.font.color.rgb = RGBColor(46, 116, 181)
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.add_run("Penunjukan Langsung · Environment Staging · 3 Agustus 2026").italic = True

doc.add_heading("Ringkasan hasil eksekusi", level=1)
doc.add_paragraph(
    "Seluruh rantai Flow 1–3 berhasil diselesaikan pada satu paket yang sama. "
    "Status akhir manifest adalah POKJA_APPROVED."
)
add_table(doc, [
    ["Atribut", "Nilai"],
    ["Package ID", "01KZ2Q7CG955CAV5V2J7NE8M89"],
    ["Nama paket", "[QA] [Penunjukan Langsung] [Automation] [009] - Perbaikan Baperkam RW.13 Kel. Kecapi"],
    ["Jenis", "Non-Tender · Pekerjaan Konstruksi · Penunjukan Langsung"],
    ["Satker", "DINAS PERUMAHAN RAKYAT DAN KAWASAN PERMUKIMAN"],
    ["Status akhir", "POKJA_APPROVED"],
], [1.5, 5.0])

doc.add_heading("Alur dan fungsi masing-masing role", level=1)
add_table(doc, [
    ["Role", "Tujuan dalam alur"],
    ["PPK", "Membentuk paket dari RUP, melengkapi HPS/KBKI dan dokumen persiapan, lalu mengirim paket kepada Kepala UKPBJ."],
    ["Kepala UKPBJ", "Memastikan paket dan dokumen persiapan lengkap, menetapkan Pokja, mengunggah Nota Dinas KUPPBJ, lalu meneruskan paket kepada Pokja."],
    ["Pokja Pemilihan", "Menetapkan konfigurasi pemilihan, jadwal, dokumen, undangan, BA reviu, penyedia target, dan menyetujui Pakta Integritas secara kolektif."],
], [1.5, 5.0])

doc.add_heading("Flow 1 — Pembuatan paket oleh PPK", level=1)
add_step(doc, "Login sebagai PPK", "Masuk ke SPSE menggunakan persona PPK pada instansi/satker yang sesuai.")
add_step(doc, "Buka pembuatan Non-Tender", "Pilih satker, tahun 2026, metode Penunjukan Langsung, dan jenis Pekerjaan Konstruksi.")
add_step(doc, "Pilih RUP", "Paket menggunakan RUP Perbaikan Baperkam RW.13 Kel. Kecapi.")
add_step(doc, "Lengkapi paket", "Isi nama paket, rincian HPS/KBKI, dan unggah seluruh dokumen persiapan.")
add_step(doc, "Kirim ke Kepala UKPBJ", "Pilih pejabat tujuan, unggah Nota Dinas PPK, simpan, dan bentuk manifest berstatus SENT_TO_UKPBJ.")

doc.add_heading("Flow 2 — Verifikasi Kepala UKPBJ", level=1)
add_step(doc, "Login sebagai Kepala UKPBJ", "Gunakan persona Kepala UKPBJ Kota Cirebon.")
add_step(doc, "Buka paket dari manifest", "Cari berdasarkan package ID dan nama exact, bukan urutan tabel.")
add_step(doc, "Verifikasi data", "Cocokkan nama paket, satker, RUP, metode, dan kelengkapan Dokumen Persiapan.")
add_step(doc, "Tentukan Pokja", "Pilih Pokja Pemilihan yang akan memproses paket.")
add_step(doc, "Unggah dan teruskan", "Unggah Nota Dinas KUPPBJ, kirim email kepada Pokja, lalu simpan hingga manifest menjadi SENT_TO_POKJA.")

doc.add_heading("Flow 3 — Persetujuan tiga anggota Pokja", level=1)
add_step(doc, "Pokja 1 login", "Konfigurasi paket dikerjakan pada sesi terisolasi.")
add_step(doc, "Atur jadwal", "Tanggal jadwal paket wajib dihitung dari tanggal aktif Sesi Default. Saat eksekusi, Sesi Default adalah 15 Agustus 2026; jadwal dibuat H+1 sampai H+9.")
add_step(doc, "Lengkapi dokumen pemilihan", "Atur kualifikasi usaha, persyaratan kualifikasi, masa berlaku penawaran, dokumen penawaran, dan unggah SDP.")
add_step(doc, "Kirim undangan", "Tanggal undangan mengikuti real date. Tidak ada validasi tanggal undangan seperti pada jadwal paket.")
add_step(doc, "BA reviu dan penyedia", "Unggah BA hasil reviu dan pilih penyedia target.")
add_step(doc, "Persetujuan kolektif", "Pokja 1, Pokja 2, dan Pokja 3 menyetujui Pakta Integritas dalam BrowserContext terpisah.")

doc.add_heading("Screenshot bukti eksekusi", level=1)
doc.add_heading("Daftar paket pada sesi Pokja", level=2)
doc.add_paragraph("Tampilan sesi anggota Pokja setelah autentikasi dan pemilihan role.")
if IMG1.exists():
    doc.add_picture(str(IMG1), width=Inches(6.4))
doc.add_heading("Paket selesai dikonfigurasi", level=2)
doc.add_paragraph("Jadwal berstatus tersimpan dan seluruh komponen Dokumen Pemilihan telah dilengkapi.")
if IMG2.exists():
    doc.add_picture(str(IMG2), width=Inches(6.4))

doc.add_heading("Catatan kendala dan perbaikan", level=1)
for text in (
    "Jadwal awal ditolak karena automation memakai tanggal riil 3 Agustus 2026, sedangkan Sesi Default berada pada 15 Agustus 2026.",
    "Automation diperbaiki agar membaca tanggal aktif Sesi Default dari /spse/training/beranda sebelum menghitung H+1 sampai H+9.",
    "Tanggal undangan tetap menggunakan real date sesuai klarifikasi dan tidak menjadi sumber validasi jadwal.",
    "Gangguan komunikasi saat deployment ditangani dengan resume; status anggota yang sudah sukses tidak diproses ulang.",
    "Redirect Security Key pascadeploy dibuat lebih tahan terhadap navigasi lambat.",
):
    doc.add_paragraph(text, style="List Bullet")

doc.add_page_break()
doc.add_heading("Status persetujuan", level=1)
add_table(doc, [
    ["Tahap", "Hasil", "Waktu (UTC)"],
    ["PPK membuat paket", "Sukses", "2026-08-03 02:30"],
    ["Kepala UKPBJ meneruskan", "Sukses", "2026-08-03 02:30"],
    ["Pokja 1", "Disetujui", "2026-08-03 02:45"],
    ["Pokja 2", "Disetujui", "2026-08-03 03:08"],
    ["Pokja 3", "Disetujui", "2026-08-03 03:09"],
], [2.2, 1.6, 2.7])

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(OUT)
