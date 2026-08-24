import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
GUIDE_ROOT = ROOT / "docs" / "visual-guide-2026-08-03"
OUT = ROOT / "docs" / "SPSE_Visual_Click_Input_Guide_Flow_1_2_3.docx"


def latest_successful_run(log_path):
    records = [
        json.loads(line)
        for line in log_path.read_text().splitlines()
        if line.strip()
    ]
    starts = [index for index, item in enumerate(records) if item["step"] == 1]
    if not starts:
        return records
    run = records[starts[-1]:]
    seen = set()
    clean = []
    for item in run:
        image_path = GUIDE_ROOT / item["image"]
        key = (item["step"], item["instruction"], item["image"])
        if key in seen or not image_path.exists():
            continue
        seen.add(key)
        clean.append(item)
    return clean


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


flows = [
    ("Flow 1 — PPK membuat paket", "flow1-ppk.jsonl"),
    ("Flow 2 — Kepala UKPBJ memverifikasi paket", "flow2-ukpbj.jsonl"),
    ("Flow 3 — Pokja mengonfigurasi dan menyetujui", "flow3-pokja.jsonl"),
]

doc = Document()
section = doc.sections[0]
section.orientation = WD_ORIENT.LANDSCAPE
section.page_width = Inches(11)
section.page_height = Inches(8.5)
section.top_margin = Inches(0.55)
section.bottom_margin = Inches(0.55)
section.left_margin = Inches(0.65)
section.right_margin = Inches(0.65)
section.header_distance = Inches(0.3)
section.footer_distance = Inches(0.3)

styles = doc.styles
styles["Normal"].font.name = "Calibri"
styles["Normal"].font.size = Pt(10.5)
styles["Normal"].paragraph_format.space_after = Pt(4)
styles["Normal"].paragraph_format.line_spacing = 1.15
for name, size, before, after in (
    ("Heading 1", 20, 10, 6),
    ("Heading 2", 15, 8, 4),
    ("Heading 3", 12, 5, 3),
):
    style = styles[name]
    style.font.name = "Calibri"
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(46, 116, 181)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
title.paragraph_format.space_before = Pt(70)
run = title.add_run("Panduan Visual SPSE\nFlow 1, Flow 2, dan Flow 3")
run.bold = True
run.font.size = Pt(28)
run.font.color.rgb = RGBColor(46, 116, 181)
subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.add_run(
    "Panduan klik, input, pilihan, upload, dan hasil yang diharapkan\n"
    "Penunjukan Langsung · Environment Staging · 3 Agustus 2026"
).italic = True
doc.add_paragraph()
lead = doc.add_paragraph()
lead.alignment = WD_ALIGN_PARAGRAPH.CENTER
lead.add_run(
    "Paket dokumentasi: Automation [011] · Package ID 01KZ2TCGG2EMMBNKP23VJG6M7V · Status akhir POKJA_APPROVED"
).bold = True

doc.add_page_break()
doc.add_heading("Cara membaca panduan", level=1)
for text in (
    "Lingkaran oranye menunjukkan tombol, field, dropdown, checkbox, atau link yang harus digunakan.",
    "Nomor pada lingkaran sama dengan nomor langkah pada halaman.",
    "Caption gelap pada screenshot menunjukkan jenis tindakan: KLIK, ISI, PILIH, atau UNGGAH.",
    "Bagian Hasil yang diharapkan dipakai sebagai checkpoint sebelum melanjutkan.",
    "Email dan sandi tidak ditampilkan dalam screenshot setelah terisi.",
):
    doc.add_paragraph(text, style="List Bullet")

doc.add_heading("Peran dalam alur", level=1)
table = doc.add_table(rows=1, cols=2)
table.style = "Table Grid"
table.autofit = False
table.columns[0].width = Inches(2.0)
table.columns[1].width = Inches(7.5)
headers = table.rows[0].cells
headers[0].text = "Role"
headers[1].text = "Fungsi"
set_repeat_table_header(table.rows[0])
for cell in headers:
    for paragraph in cell.paragraphs:
        for cell_run in paragraph.runs:
            cell_run.bold = True
for role, purpose in (
    ("PPK", "Membentuk paket dari RUP, melengkapi HPS/KBKI dan dokumen persiapan, lalu mengirim paket kepada Kepala UKPBJ."),
    ("Kepala UKPBJ", "Memverifikasi paket, menetapkan Pokja, mengunggah Nota Dinas KUPPBJ, lalu meneruskan paket."),
    ("Pokja", "Mengatur jadwal dan dokumen pemilihan, mengirim undangan, memilih penyedia, lalu menyetujui Pakta Integritas."),
):
    cells = table.add_row().cells
    cells[0].text = role
    cells[1].text = purpose

total_steps = 0
for flow_title, log_name in flows:
    records = latest_successful_run(GUIDE_ROOT / log_name)
    doc.add_page_break()
    heading = doc.add_heading(flow_title, level=1)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    intro = doc.add_paragraph()
    intro.alignment = WD_ALIGN_PARAGRAPH.CENTER
    intro.add_run(f"{len(records)} langkah visual").italic = True

    for index, item in enumerate(records, start=1):
        doc.add_page_break()
        total_steps += 1
        step_title = doc.add_paragraph()
        step_title.paragraph_format.space_after = Pt(3)
        number_run = step_title.add_run(f"Langkah {index} · {item['action'].upper()} — ")
        number_run.bold = True
        title_size = 13 if len(item["instruction"]) > 34 else 16
        number_run.font.size = Pt(title_size)
        number_run.font.color.rgb = RGBColor(46, 116, 181)
        text_run = step_title.add_run(item["instruction"])
        text_run.bold = True
        text_run.font.size = Pt(title_size)

        expected = doc.add_paragraph()
        expected.paragraph_format.space_after = Pt(5)
        expected.add_run("Hasil yang diharapkan: ").bold = True
        expected.add_run(item.get("expected") or "Tampilan berpindah ke tahap berikutnya tanpa error.")

        image_path = GUIDE_ROOT / item["image"]
        doc.add_picture(str(image_path), width=Inches(9.65))
        picture_paragraph = doc.paragraphs[-1]
        picture_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        picture_paragraph.paragraph_format.space_after = Pt(0)

footer = section.footer
footer_p = footer.paragraphs[0]
footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
footer_p.add_run(
    f"SPSE Visual Guide · {total_steps} langkah · Package 01KZ2TCGG2EMMBNKP23VJG6M7V"
).font.size = Pt(8)

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
print(json.dumps({"output": str(OUT), "steps": total_steps}))
